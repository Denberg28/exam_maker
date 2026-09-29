import io
import json
import tempfile
import threading
import unittest
from pathlib import Path

from openpyxl import load_workbook
from web.core import initialize, create_session, start_attempt, attempt_view, submit_answer, results, export_xlsx, set_session_active, question_template, bank_from_csv

BANK = {"schema": 1, "title": "Sample", "questions": [
    {"id": "a", "prompt": "One?", "options": ["A", "B", "C", "D"], "correct": 1},
    {"id": "b", "prompt": "Two?", "options": ["W", "X", "Y", "Z"], "correct": 2},
]}

class SessionTests(unittest.TestCase):
    def test_csv_template_import(self):
        raw = question_template() + '"What is a rivet, used for?",Join,Separate,Paint,Measure,A,\r\n'.encode()
        bank = bank_from_csv(raw, "Structures")
        self.assertEqual("What is a rivet, used for?", bank["questions"][0]["prompt"])
        self.assertEqual(0, bank["questions"][0]["correct"])
        with self.assertRaisesRegex(ValueError, "Row 2"):
            bank_from_csv(question_template() + b"Question,A,B,C,D,E,\r\n", "Structures")
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.db = Path(self.tmp.name) / "exam.sqlite"
        initialize(self.db)
    def tearDown(self):
        self.tmp.cleanup()
    def test_two_sessions_and_concurrent_examiners_are_isolated(self):
        first = create_session(self.db, BANK, "Morning")
        second = create_session(self.db, BANK, "Afternoon")
        set_session_active(self.db, first, True)
        set_session_active(self.db, second, True)
        ids = []
        def join(code, name):
            ids.append((code, start_attempt(self.db, code, name, name + "-ID")))
        threads = [threading.Thread(target=join, args=(first, "Ada")), threading.Thread(target=join, args=(second, "Ben"))]
        for t in threads: t.start()
        for t in threads: t.join()
        self.assertEqual(2, len(ids))
        for code, aid in ids:
            self.assertEqual(code, attempt_view(self.db, aid)["session_code"])
            for _ in range(2):
                view = attempt_view(self.db, aid)
                self.assertNotIn("correct", view["question"])
                submit_answer(self.db, aid, view["index"], 0)
            self.assertTrue(attempt_view(self.db, aid)["completed"])
            with self.assertRaises(ValueError): submit_answer(self.db, aid, 0, 0)
        self.assertEqual(1, len(results(self.db, first)))
        self.assertEqual(1, len(results(self.db, second)))
        wb = load_workbook(io.BytesIO(export_xlsx(results(self.db))))
        self.assertEqual(3, wb.active.max_row)
    def test_closed_session_blocks_new_attempt_but_existing_can_finish(self):
        code = create_session(self.db, BANK, "Morning")
        set_session_active(self.db, code, True)
        aid = start_attempt(self.db, code, "A", "1")
        set_session_active(self.db, code, False)
        with self.assertRaises(ValueError): start_attempt(self.db, code, "B", "2")
        self.assertEqual(2, attempt_view(self.db, aid)["score"])
        submit_answer(self.db, aid, 0, 0)
    def test_draft_requires_publish_and_closed_session_keeps_results(self):
        code = create_session(self.db, BANK, "Trial")
        with self.assertRaises(ValueError): start_attempt(self.db, code, "A", "1")
        set_session_active(self.db, code, True)
        aid = start_attempt(self.db, code, "A", "1")
        set_session_active(self.db, code, False)
        for _ in range(2):
            view = attempt_view(self.db, aid)
            submit_answer(self.db, aid, view["index"], 0)
        self.assertEqual(1, len(results(self.db, code)))
    def test_restart_preserves_score_and_duplicate_submission_is_rejected(self):
        code = create_session(self.db, BANK, "Morning")
        set_session_active(self.db, code, True)
        aid = start_attempt(self.db, code, "A", "1")
        first = attempt_view(self.db, aid)
        submit_answer(self.db, aid, first["index"], 0)
        initialize(self.db)  # a new app process opens the existing database
        resumed = attempt_view(self.db, aid)
        self.assertEqual(1, resumed["index"])
        with self.assertRaises(ValueError):
            submit_answer(self.db, aid, first["index"], 1)
        self.assertEqual(resumed["score"], attempt_view(self.db, aid)["score"])
        submit_answer(self.db, aid, resumed["index"], 0)
        self.assertEqual(1, len(results(self.db, code)))
    def test_rejects_bad_bank(self):
        bad = json.loads(json.dumps(BANK)); bad["questions"][0]["options"] = ["A", "B"]
        with self.assertRaises(ValueError): create_session(self.db, bad, "Bad")
    def test_excel_formula_is_text(self):
        data = export_xlsx([("now", "=HYPERLINK(\"x\")", "-1", "S", 1, 1, 100)])
        cells = load_workbook(io.BytesIO(data)).active
        self.assertEqual('=HYPERLINK("x")', cells["B2"].value[1:])
        self.assertEqual("s", cells["B2"].data_type)
